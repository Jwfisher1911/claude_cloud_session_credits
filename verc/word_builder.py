"""Turn a VERC Markdown draft into a formatted Word document.

Supported Markdown (the same rules Claude is told to follow):
  # Title                 skipped (the cover block already shows the title)
  ## Heading / ### / ####  Word headings 1, 2, 3
  - item / * item          bullets (indent two spaces for a second level)
  1. item                  numbered steps
  - [ ] item               checkbox line
  > text                   shaded callout box
  | a | b |                table (empty cells become writing space)
  [[LINES 3]]              three ruled lines for handwritten answers
  [[PAGEBREAK]]            new page
  **bold**, *italic*, [VERIFY: ...] (highlighted yellow)
"""

from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_COLOR_INDEX, WD_PARAGRAPH_ALIGNMENT, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\[VERIFY[^\]]*\]|(?<!\*)\*[^*\s][^*]*?\*(?!\*))", re.IGNORECASE)
LINES_RE = re.compile(r"^\s*\[\[LINES\s+(\d+)\]\]\s*$", re.IGNORECASE)
PAGEBREAK_RE = re.compile(r"^\s*\[\[PAGEBREAK\]\]\s*$", re.IGNORECASE)
CHECK_RE = re.compile(r"^(\s*)[-*]\s+\[[ xX]?\]\s*(.*)$")
BULLET_RE = re.compile(r"^(\s*)[-*•]\s+(.*)$")
NUMBER_RE = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")
HEADING_RE = re.compile(r"^(#{1,4})\s+(.*?)\s*#*\s*$")
TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


@dataclass
class DocInfo:
    program_name: str
    doc_type: str  # "Lesson Guide" or "Workbook"
    title: str
    subtitle: str  # e.g. "Module 1 of 7 | VERC Core Modules"
    goal: str = ""
    disclaimer: str = ""
    legal_line: str = ""
    website: str = ""
    colors: dict = field(default_factory=dict)
    font: str = "Calibri"
    break_before_h2: bool = False
    release: bool = False  # True for buyer copies: no draft banner


def _hex(color: str, fallback: str) -> RGBColor:
    value = (color or fallback).lstrip("#")
    return RGBColor.from_string(value.upper())


def _shade(element, fill_hex: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex.lstrip("#"))
    element.append(shd)


def _paragraph_border(paragraph, side: str, color_hex: str, size: int = 8, space: int = 1) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    edge = OxmlElement(f"w:{side}")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size))
    edge.set(qn("w:space"), str(space))
    edge.set(qn("w:color"), color_hex.lstrip("#"))
    borders.append(edge)


def _add_field(paragraph, instruction: str, size_half_points: int = 16) -> None:
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instruction)
    inner = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(size_half_points))
    r_pr.append(sz)
    inner.append(r_pr)
    text = OxmlElement("w:t")
    text.text = "1"
    inner.append(text)
    fld.append(inner)
    paragraph._p.append(fld)


def _set_font(style, name: str) -> None:
    style.font.name = name
    r_pr = style.element.get_or_add_rPr()
    fonts = r_pr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if fonts.get(qn(attr)) is not None:
            del fonts.attrib[qn(attr)]


def add_inline(paragraph, text: str, bold: bool = False, color: RGBColor | None = None, size: Pt | None = None) -> None:
    """Add text with **bold**, *italic*, and highlighted [VERIFY] markers."""
    text = text.replace("`", "")
    for piece in INLINE_RE.split(text):
        if not piece:
            continue
        run_bold, run_italic, highlight = bold, False, False
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            piece, run_bold = piece[2:-2], True
        elif piece.upper().startswith("[VERIFY"):
            highlight, run_bold = True, True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            piece, run_italic = piece[1:-1], True
        run = paragraph.add_run(piece)
        run.bold = run_bold or None
        run.italic = run_italic or None
        if highlight:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if color is not None:
            run.font.color.rgb = color
        if size is not None:
            run.font.size = size


class WordBuilder:
    def __init__(self, info: DocInfo):
        self.info = info
        self.primary = _hex(info.colors.get("primary", ""), "1F3A5F")
        self.accent = _hex(info.colors.get("accent", ""), "B8862B")
        self.primary_hex = str(self.primary)
        self.accent_hex = str(self.accent)
        self.callout_fill = info.colors.get("callout_fill", "EEF2F7")
        self.doc = Document()
        self._seen_h2 = False
        self._setup_page()

    # ----- document frame -------------------------------------------------
    def _setup_page(self) -> None:
        section = self.doc.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        for side in ("left_margin", "right_margin"):
            setattr(section, side, Inches(1))
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)

        styles = self.doc.styles
        normal = styles["Normal"]
        _set_font(normal, self.info.font)
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15

        for name, size, space_before in (("Heading 1", 16, 18), ("Heading 2", 13, 12), ("Heading 3", 11.5, 10)):
            style = styles[name]
            _set_font(style, self.info.font)
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = self.primary
            style.paragraph_format.space_before = Pt(space_before)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.keep_with_next = True

        for name in ("List Bullet", "List Bullet 2"):
            _set_font(styles[name], self.info.font)
            styles[name].paragraph_format.space_after = Pt(3)

        header = section.header.paragraphs[0]
        header.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
        add_inline(header, f"VERC  |  {self.info.title}", color=RGBColor(0x6B, 0x72, 0x80), size=Pt(8.5))

        footer = section.footer.paragraphs[0]
        footer.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        disclaimer = " ".join(x for x in (self.info.disclaimer, self.info.legal_line) if x)
        add_inline(footer, disclaimer, color=RGBColor(0x6B, 0x72, 0x80), size=Pt(7.5))
        page = section.footer.add_paragraph()
        page.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        run = page.add_run("Page ")
        run.font.size = Pt(8)
        _add_field(page, "PAGE")

    def _cover_block(self) -> None:
        info = self.info
        p = self.doc.add_paragraph()
        add_inline(p, info.program_name.upper(), bold=True, color=self.accent, size=Pt(9))
        p.paragraph_format.space_after = Pt(2)

        p = self.doc.add_paragraph()
        add_inline(p, f"{info.doc_type.upper()}  |  {info.subtitle}", color=RGBColor(0x6B, 0x72, 0x80), size=Pt(9))
        p.paragraph_format.space_after = Pt(10)

        p = self.doc.add_paragraph()
        add_inline(p, info.title, bold=True, color=self.primary, size=Pt(24))
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.0

        if info.goal:
            p = self.doc.add_paragraph()
            add_inline(p, f"*{info.goal}*", size=Pt(12))
            p.paragraph_format.space_after = Pt(8)

        p = self.doc.add_paragraph()
        stamp = dt.date.today().strftime("%B %d, %Y").replace(" 0", " ")
        if info.release:
            add_inline(p, f"Current as of {dt.date.today().strftime('%B %Y')}.", size=Pt(8.5))
        else:
            add_inline(p, f"DRAFT generated {stamp}. Review every fact and every highlighted VERIFY note before publishing.",
                       color=RGBColor(0x9A, 0x3B, 0x2E), size=Pt(8.5))
        p.paragraph_format.space_after = Pt(14)
        _paragraph_border(p, "bottom", self.accent_hex, size=12, space=6)

    # ----- block renderers ------------------------------------------------
    def _heading(self, level: int, text: str) -> None:
        if level == 1:
            return  # The cover block already carries the title.
        if level == 2 and self.info.break_before_h2:
            if self._seen_h2:
                self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            self._seen_h2 = True
        style = {2: "Heading 1", 3: "Heading 2"}.get(level, "Heading 3")
        para = self.doc.add_paragraph(style=style)
        add_inline(para, text)

    def _paragraph(self, text: str) -> None:
        add_inline(self.doc.add_paragraph(), text)

    def _bullet(self, text: str, level: int) -> None:
        para = self.doc.add_paragraph(style="List Bullet 2" if level > 0 else "List Bullet")
        add_inline(para, text)

    def _numbered(self, number: str, text: str, level: int) -> None:
        para = self.doc.add_paragraph()
        fmt = para.paragraph_format
        fmt.left_indent = Inches(0.35 + 0.3 * level)
        fmt.first_line_indent = Inches(-0.3)
        fmt.space_after = Pt(3)
        run = para.add_run(f"{number}.\t")
        run.bold = True
        run.font.color.rgb = self.primary
        fmt.tab_stops.add_tab_stop(Inches(0.35 + 0.3 * level))
        add_inline(para, text)

    def _checkbox(self, text: str, level: int) -> None:
        para = self.doc.add_paragraph()
        fmt = para.paragraph_format
        fmt.left_indent = Inches(0.35 + 0.3 * level)
        fmt.first_line_indent = Inches(-0.3)
        fmt.space_after = Pt(4)
        fmt.tab_stops.add_tab_stop(Inches(0.35 + 0.3 * level))
        box = para.add_run("☐\t")
        box.font.size = Pt(13)
        box.font.color.rgb = self.primary
        add_inline(para, text)

    def _callout(self, lines: list[str]) -> None:
        para = self.doc.add_paragraph()
        fmt = para.paragraph_format
        fmt.left_indent = Inches(0.15)
        fmt.right_indent = Inches(0.15)
        fmt.space_before = Pt(6)
        fmt.space_after = Pt(10)
        _shade(para._p.get_or_add_pPr(), self.callout_fill)
        _paragraph_border(para, "left", self.accent_hex, size=24, space=8)
        for i, line in enumerate(lines):
            if i:
                para.add_run().add_break()
            add_inline(para, line)

    def _writing_lines(self, count: int) -> None:
        # A right tab with an underscore leader draws a full-width ruled line.
        for _ in range(max(1, min(count, 20))):
            para = self.doc.add_paragraph()
            fmt = para.paragraph_format
            fmt.space_before = Pt(12)
            fmt.space_after = Pt(0)
            fmt.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.LINES)
            run = para.add_run("\t")
            run.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def _table(self, rows: list[list[str]]) -> None:
        if not rows:
            return
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
        table = self.doc.add_table(rows=len(rows), cols=width)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = self._column_widths(rows)
        for c_index, col in enumerate(table.columns):
            col.width = widths[c_index]
        for r_index, row in enumerate(rows):
            tr = table.rows[r_index]
            is_header = r_index == 0
            is_example = not is_header and row[0].strip().lower().startswith("example")
            if is_header:
                tr_pr = tr._tr.get_or_add_trPr()
                repeat = OxmlElement("w:tblHeader")
                repeat.set(qn("w:val"), "true")
                tr_pr.append(repeat)
            elif not any(cell.strip() for cell in row):
                tr.height = Inches(0.38)
                tr.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            for c_index, value in enumerate(row):
                cell = tr.cells[c_index]
                cell.width = widths[c_index]
                para = cell.paragraphs[0]
                para.paragraph_format.space_after = Pt(2)
                if is_header:
                    _shade(cell._tc.get_or_add_tcPr(), self.primary_hex)
                    add_inline(para, value.strip(), bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=Pt(10))
                elif is_example:
                    add_inline(para, f"*{value.strip()}*" if value.strip() else "", color=RGBColor(0x6B, 0x72, 0x80), size=Pt(10))
                else:
                    add_inline(para, value.strip(), size=Pt(10))
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    @staticmethod
    def _column_widths(rows: list[list[str]], total: float = 6.5) -> list:
        """Size columns by content. Columns left blank for writing get extra room."""
        weights = []
        body_rows = [r for r in rows[1:] if not r[0].strip().lower().startswith("example")]
        for c in range(len(rows[0])):
            header = rows[0][c].strip()
            body = [r[c].strip() for r in body_rows if r[c].strip()]
            longest = max([len(header)] + [len(x) for x in body])
            weight = min(max(longest, 3), 40)
            if not body:  # a column the reader fills in by hand
                weight = max(weight, 14)
            # Never narrower than the longest single word, so "Wednesday" stays on one line.
            words = [w for r in rows for w in r[c].split()]
            weight = max(weight, (max((len(w) for w in words), default=0) + 3) * 1.25)
            weights.append(weight)
        scale = total / sum(weights)
        return [Inches(w * scale) for w in weights]

    # ----- parser ---------------------------------------------------------
    @staticmethod
    def _split_row(line: str) -> list[str]:
        line = line.strip()
        if line.startswith("|"):
            line = line[1:]
        if line.endswith("|"):
            line = line[:-1]
        return [cell.strip() for cell in line.split("|")]

    def build(self, markdown: str) -> Document:
        self._cover_block()
        lines = markdown.replace("\r\n", "\n").split("\n")
        i, para_buffer = 0, []

        def flush() -> None:
            if para_buffer:
                self._paragraph(" ".join(s.strip() for s in para_buffer))
                para_buffer.clear()

        while i < len(lines):
            line = lines[i]
            stripped = line.strip()
            indent = (len(line) - len(line.lstrip(" "))) // 2

            if not stripped or stripped in ("---", "***", "___"):
                flush()
                i += 1
                continue
            if stripped.startswith("```"):
                flush()
                i += 1
                continue
            if m := HEADING_RE.match(stripped):
                flush()
                self._heading(len(m.group(1)), m.group(2))
            elif m := LINES_RE.match(stripped):
                flush()
                self._writing_lines(int(m.group(1)))
            elif PAGEBREAK_RE.match(stripped):
                flush()
                self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            elif stripped.startswith("|"):
                flush()
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    if not TABLE_SEP_RE.match(lines[i]):
                        rows.append(self._split_row(lines[i]))
                    i += 1
                self._table(rows)
                continue
            elif stripped.startswith(">"):
                flush()
                quote = []
                while i < len(lines) and lines[i].strip().startswith(">"):
                    quote.append(lines[i].strip()[1:].strip())
                    i += 1
                self._callout([q for q in quote if q] or [""])
                continue
            elif m := CHECK_RE.match(line):
                flush()
                self._checkbox(m.group(2), indent)
            elif m := NUMBER_RE.match(line):
                flush()
                self._numbered(m.group(2), m.group(3), indent)
            elif m := BULLET_RE.match(line):
                flush()
                self._bullet(m.group(2), indent)
            else:
                para_buffer.append(stripped)
            i += 1
        flush()
        return self.doc


def markdown_to_docx(markdown: str, info: DocInfo, out_path: Path) -> Path:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    WordBuilder(info).build(markdown).save(str(out_path))
    return out_path
